import io

import pytest
from docx import Document
from pypdf import PdfWriter

from surv4impact.parser import (
    QuestionnaireParseError,
    _assemble_questionnaire,
    _TableBlock,
    _TextBlock,
    parse_docx,
    parse_pdf,
    parse_questionnaire,
)


def sample_docx() -> bytes:
    document = Document()
    document.add_heading("Experiência", level=1)
    document.add_paragraph("1. Nos últimos 12 meses, utilizou o serviço?")
    document.add_paragraph("☐ Sim")
    document.add_paragraph("☐ Não")
    document.add_paragraph("2. Como avalia o apoio recebido?")
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "3. Avalie as dimensões seguintes:"
    table.cell(0, 1).text = "Baixo"
    table.cell(0, 2).text = "Elevado"
    table.cell(1, 0).text = "Clareza"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_parser_extracts_questions_options_and_matrix():
    questionnaire = parse_docx(sample_docx(), "teste.docx")

    assert questionnaire.title == "teste"
    assert len(questionnaire.questions) == 3
    assert questionnaire.questions[0].options == ["Sim", "Não"]
    assert questionnaire.questions[0].question_type == "sim/não"
    assert questionnaire.questions[2].matrix_columns == 3
    assert "Experiência" in questionnaire.sections


def test_parser_does_not_mistake_uppercase_checkbox_options_for_headings():
    document = Document()
    document.add_paragraph("1. Utilizou o serviço?")
    document.add_paragraph("☐ SIM")
    document.add_paragraph("☐ NÃO")
    buffer = io.BytesIO()
    document.save(buffer)

    questionnaire = parse_docx(buffer.getvalue(), "opcoes.docx")

    assert questionnaire.questions[0].options == ["SIM", "NÃO"]
    assert questionnaire.questions[0].question_type == "sim/não"


def test_pdf_parser_extracts_title_sections_questions_options_and_matrix(
    questionnaire_pdf_bytes: bytes,
):
    questionnaire = parse_pdf(questionnaire_pdf_bytes, "questionario.pdf")

    assert questionnaire.title == "Questionário PDF de teste"
    assert questionnaire.source_name == "questionario.pdf"
    assert len(questionnaire.questions) == 3
    assert questionnaire.questions[0].options == ["Sim", "Não"]
    assert questionnaire.questions[0].question_type == "sim/não"
    assert questionnaire.questions[2].matrix_rows == 1
    assert questionnaire.questions[2].matrix_columns == 3
    assert "EXPERIÊNCIA" in questionnaire.sections
    assert "Página 1" in questionnaire.questions[0].source
    assert "Página 2" in questionnaire.questions[2].source
    assert "Como avalia o apoio recebido?" in questionnaire.raw_text


def test_pdf_parser_warns_when_only_some_pages_have_no_text(
    mixed_blank_pdf_bytes: bytes,
):
    questionnaire = parse_pdf(mixed_blank_pdf_bytes, "pagina-vazia.pdf")

    assert questionnaire.questions
    blank_page_warning = next(
        warning
        for warning in questionnaire.parse_warnings
        if "texto" in warning.casefold()
    )
    assert "página" in blank_page_warning.casefold()
    assert "1" in blank_page_warning


def test_pdf_parser_rejects_pdf_without_extractable_text(empty_pdf_bytes: bytes):
    with pytest.raises(QuestionnaireParseError):
        parse_pdf(empty_pdf_bytes, "digitalizado.pdf")


def test_pdf_parser_rejects_corrupt_pdf():
    with pytest.raises(QuestionnaireParseError):
        parse_pdf(b"%PDF-1.4\nficheiro truncado", "corrompido.pdf")


def test_pdf_parser_accepts_a_valid_header_with_a_short_prefix(
    questionnaire_pdf_bytes: bytes,
):
    questionnaire = parse_pdf(
        b"\xef\xbb\xbf\r\n" + questionnaire_pdf_bytes,
        "com-prefixo.pdf",
    )

    assert len(questionnaire.questions) == 3


def test_pdf_parser_keeps_repeated_options_near_page_tops(
    repeated_options_pdf_bytes: bytes,
):
    questionnaire = parse_pdf(repeated_options_pdf_bytes, "opcoes-repetidas.pdf")

    assert len(questionnaire.questions) == 2
    assert [question.options for question in questionnaire.questions] == [
        ["Sim", "Não"],
        ["Sim", "Não"],
    ]


def test_pdf_parser_merges_a_wrapped_numbered_prompt_after_a_period(
    wrapped_question_pdf_bytes: bytes,
):
    questionnaire = parse_pdf(wrapped_question_pdf_bytes, "pergunta-repartida.pdf")

    assert len(questionnaire.questions) == 1
    assert questionnaire.questions[0].text == (
        "Recorde a sua participação no programa. "
        "Que motivo melhor explica a sua escolha?"
    )


@pytest.mark.parametrize(
    "rows",
    [
        pytest.param([["Sim", "Não"]], id="single-row"),
        pytest.param([["Sim"], ["Não"]], id="single-column"),
        pytest.param([["", "Sim"], ["", "Não"]], id="blank-stub-column"),
    ],
)
def test_option_only_table_shapes_attach_options_without_creating_a_matrix(rows):
    questionnaire = _assemble_questionnaire(
        [
            _TextBlock(text="1. Aceita participar?", source="Página 1, linha 1"),
            _TableBlock(rows=rows, source="Página 1, tabela 1"),
        ],
        title="Opções em tabela",
        filename="opcoes.pdf",
        format_name="PDF",
    )

    assert len(questionnaire.questions) == 1
    assert questionnaire.questions[0].options == ["Sim", "Não"]
    assert questionnaire.questions[0].matrix_rows == 0
    assert questionnaire.questions[0].question_type == "sim/não"


def test_inline_pdf_options_are_split_into_distinct_answers():
    questionnaire = _assemble_questionnaire(
        [
            _TextBlock(text="1. Aceita participar?", source="Página 1, linha 1"),
            _TextBlock(text="a) Sim    b) Não", source="Página 1, linha 2"),
        ],
        title="Opções horizontais",
        filename="opcoes.pdf",
        format_name="PDF",
    )

    assert questionnaire.questions[0].options == ["Sim", "Não"]
    assert questionnaire.questions[0].question_type == "sim/não"


def test_pdf_parser_keeps_a_bare_year_in_the_middle_of_a_page(
    mid_page_year_pdf_bytes: bytes,
):
    questionnaire = parse_pdf(mid_page_year_pdf_bytes, "ano.pdf")

    assert "2024" in questionnaire.raw_text


def test_pdf_parser_rejects_text_document_without_questions(
    text_only_pdf_bytes: bytes,
):
    with pytest.raises(QuestionnaireParseError):
        parse_pdf(text_only_pdf_bytes, "sem-perguntas.pdf")


def test_pdf_parser_rejects_password_protected_pdf():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("segredo")
    buffer = io.BytesIO()
    writer.write(buffer)

    with pytest.raises(QuestionnaireParseError, match="palavra-passe"):
        parse_pdf(buffer.getvalue(), "protegido.pdf")


def test_pdf_parser_rejects_pdf_without_pages():
    writer = PdfWriter()
    buffer = io.BytesIO()
    writer.write(buffer)

    with pytest.raises(QuestionnaireParseError, match="não contém páginas"):
        parse_pdf(buffer.getvalue(), "sem-paginas.pdf")


def test_parse_questionnaire_dispatches_docx_and_pdf(questionnaire_pdf_bytes: bytes):
    docx_questionnaire = parse_questionnaire(sample_docx(), "TESTE.DOCX")
    pdf_questionnaire = parse_questionnaire(questionnaire_pdf_bytes, "TESTE.PDF")

    assert len(docx_questionnaire.questions) == 3
    assert len(pdf_questionnaire.questions) == 3


def test_parse_questionnaire_rejects_unsupported_file_type():
    with pytest.raises(QuestionnaireParseError):
        parse_questionnaire("conteúdo".encode(), "questionario.txt")


def test_parse_questionnaire_rejects_empty_file():
    with pytest.raises(QuestionnaireParseError, match="vazio"):
        parse_questionnaire(b"", "questionario.pdf")
