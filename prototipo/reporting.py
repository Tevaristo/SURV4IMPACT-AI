"""Reutilização do padrão DOCX/JSON em memória; conteúdo inteiramente novo."""
import io
import json
from copy import deepcopy
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .norma import CRITERIOS, ROOT
from .state import completed, now

LIMITES = "Apreciação exclusivamente documental de D2.3, D2.4 e D2.6. Não demonstra compreensão efetiva, ausência empírica de enviesamentos ou funcionamento da programação. Sem classificação global. Casos localizados não validam integralmente os critérios."


def readable(value):
    if isinstance(value, bool):
        return "sim" if value else "não"
    if value is None or value == "":
        return "não registado / não pertinente"
    if isinstance(value, dict):
        return "\n".join(f"{k.replace('_', ' ').capitalize()}: {readable(v)}" for k, v in value.items())
    if isinstance(value, list):
        return "\n".join(readable(v) for v in value) or "nenhum registo"
    return str(value)


def report_data(s):
    data = deepcopy(s)
    data.update({"produto": "SURV4IMPACT AI — Protótipo v01", "gerado": now(),
        "tipo_relatorio": "Concluído" if completed(s) else "Parcial",
        "criterios_por_concluir": [c for c in CRITERIOS if c not in s["decisoes"]], "limites_gerais": LIMITES})
    return data


def export_json(s):
    return json.dumps(report_data(s), ensure_ascii=False, indent=2).encode("utf-8")


def sections(s):
    d = report_data(s)
    yield "Identificação", [f"Relatório {d['tipo_relatorio'].lower()}", f"Instrumento: {s['instrumento'].get('designacao', s['instrumento'].get('nome', 'Não identificado'))}",
        f"Versão: {s['instrumento'].get('versao', 'não identificada')}",
        "Regras: " + json.dumps(s["norma"], ensure_ascii=False),
        f"Representação: {'confirmada explicitamente' if s['confirmada'] else 'por confirmar'}; revisão {s['revisao']}",
        f"Estado técnico API: {readable(s['api'])}", LIMITES]
    yield "Contexto e proveniência", [f"{k.capitalize()}\n{readable(v)}" for k, v in s["contexto"].items()]
    yield "Documentos considerados", [f"Instrumento: {s['instrumento'].get('nome', '')}; SHA-256: {s['instrumento'].get('sha256', '')}",
        "Tipo de hash: " + s["instrumento"].get("tipo_hash", "ficheiro original ou texto inicialmente fornecido"),
        "Revisões confirmadas: " + readable(s["instrumento"].get("evidencias_revisao", [])),
        "Localização do conteúdo que substituiu a leitura automática: " + readable(
            s.get("rastreabilidade_substituicao", {}).get("localizacao_documento_original", "não aplicável")
        )] + [readable(d) for d in s["documentos"]]
    for c in CRITERIOS:
        lines = []
        p = s["propostas"].get(c)
        lines.append("Proposta: " + (p.get("categoria") or p["estado"] if p else "Por concluir"))
        if p:
            lines.extend(f"{k.replace('_', ' ')}: {v}" for k, v in p.items() if k not in ("criterio", "categoria"))
        lines.append("Decisão humana: " + (readable(s["decisoes"][c]) if c in s["decisoes"] else "Não registada"))
        yield c + " — apreciação e decisão", lines
        for rule, verification in s["verificacoes"].items():
            if rule.startswith(c):
                yield rule + " — " + verification["estado"], [verification["fundamento"]] + [
                    readable(o) for o in verification["observacoes"]]
    yield "Pedidos e respostas", [readable(p) for p in s["pedidos"]] or ["Sem pedidos registados."]
    yield "Ações e incorporação", [readable(a) for a in s["acoes"].values()] or ["Sem ações registadas."]
    yield "Representação utilizada", [readable(e) for e in s["representacao"]]
    yield "Histórico e limites de fecho", [readable({k: v for k, v in h.items() if k not in ("antes", "depois")}) for h in s["historico"]] + [
        "Critérios por concluir: " + (", ".join(d["criterios_por_concluir"]) or "nenhum"),
        "Uma correção da representação não altera o instrumento. Uma proposta aceite não confirma incorporação. O JSON conserva o texto bruto e todo o histórico da sessão."]


def export_docx(s):
    document = Document()
    sec = document.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.492)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    # standard_business_brief; override institucional: títulos #2B5B5C.
    normal.paragraph_format.space_before, normal.paragraph_format.space_after = Pt(0), Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, before, after in (("Title", 24, 0, 8), ("Heading 1", 16, 16, 8), ("Heading 2", 13, 12, 6), ("Heading 3", 12, 8, 4)):
        style = document.styles[name]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor.from_string("2B5B5C")
        style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(before), Pt(after)
        style.paragraph_format.line_spacing = 1.10
        style.paragraph_format.keep_with_next = True
    p = document.add_paragraph()
    for file, width in (("ipps-iscte.png", 1.4), ("adc.png", 1.4)):
        p.add_run().add_picture(str(ROOT / "assets/brand" / file), width=Inches(width))
        p.add_run("    ")
    document.add_paragraph().add_run().add_picture(str(ROOT / "assets/brand/pat2030-portugal2030-ue.png"), width=Inches(4.8))
    document.add_heading("SURV4IMPACT AI", 0)
    document.add_paragraph("Assistente de Revisão Metodológica de Questionários")
    for heading, paragraphs in sections(s):
        document.add_heading(heading, 1)
        for text in paragraphs:
            document.add_paragraph(str(text))
    footer = sec.footer.paragraphs[0]
    footer.add_run("SURV4IMPACT AI · ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()
