"""Exportação em memória da leitura editável do questionário."""
from csv import writer
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Element


SHEET_NAME = "Leitura do questionário"
XLSX_FILENAME = "leitura_questionario.xlsx"
CSV_FILENAME = "leitura_questionario.csv"

# A ordem e os títulos são partilhados pela tabela da interface e pelos ficheiros.
READING_COLUMNS = (
    ("id", "Referência"),
    ("texto", "Pergunta, instrução ou bloco"),
    ("localizacao", "Página ou localização no original"),
    ("formato", "Formato de resposta"),
    ("opcoes", "Opções ou escala de resposta"),
    ("subitens", "Subperguntas"),
    ("linhas", "Linhas da grelha"),
    ("colunas", "Colunas da grelha"),
    ("instrucoes", "Instruções aplicáveis"),
    ("instrucao_aplica_a", "Perguntas a que a instrução se aplica"),
    ("condicoes_destinos", "Condições de percurso e destinos"),
    ("regra_selecao", "Regra de seleção"),
    ("numero_minimo", "Número mínimo"),
    ("numero_maximo", "Número máximo"),
    ("outra_condicao_resposta", "Outra condição de resposta"),
)


def _text(value):
    """Conserva conteúdo e quebras de linha, convertendo apenas o tipo."""
    if value is None:
        return ""
    if isinstance(value, list):
        return "\n".join(str(item) for item in value)
    return str(value)


def reading_rows(elements):
    """Devolve exatamente as colunas e células mostradas no editor."""
    normalized = [Element.model_validate(element).model_dump() for element in elements]
    return [{key: _text(element.get(key, "")) for key, _ in READING_COLUMNS}
            for element in normalized]


def export_reading_xlsx(elements):
    """Cria um livro XLSX apenas com valores textuais, inteiramente em memória."""
    rows = reading_rows(elements)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    headers = [label for _, label in READING_COLUMNS]

    for column, header in enumerate(headers, 1):
        cell = sheet.cell(1, column, header)
        cell.data_type = "s"
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2B5B5C")
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    for row_number, row in enumerate(rows, 2):
        for column, (key, _) in enumerate(READING_COLUMNS, 1):
            cell = sheet.cell(row_number, column, row[key])
            # Impede que =, +, - ou @ sejam interpretados como fórmulas,
            # sem alterar o texto efetivamente guardado na célula.
            cell.data_type = "s"
            cell.number_format = "@"
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows) + 1)}"
    widths = (16, 48, 28, 24, 34, 28, 34, 34, 30, 30, 34, 32, 16, 16, 40)
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _csv_text(value):
    text = _text(value)
    # CSV não transporta tipos. O apóstrofo impede a execução quando o ficheiro
    # é aberto diretamente no Excel.
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def export_reading_csv(elements):
    """Cria CSV UTF-8 com BOM, delimitador regional e fórmulas neutralizadas."""
    output = StringIO(newline="")
    csv = writer(output, delimiter=";", lineterminator="\r\n")
    csv.writerow([label for _, label in READING_COLUMNS])
    for row in reading_rows(elements):
        csv.writerow([_csv_text(row[key]) for key, _ in READING_COLUMNS])
    return output.getvalue().encode("utf-8-sig")
