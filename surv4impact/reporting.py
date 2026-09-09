from __future__ import annotations

import io
import json
from datetime import datetime

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from . import __version__
from .knowledge import criterion_label
from .models import Finding, Questionnaire, QuestionnaireMetadata, Synthesis


def findings_dataframe(findings: list[Finding]) -> pd.DataFrame:
    rows = []
    for finding in findings:
        rows.append(
            {
                "ID": finding.finding_id,
                "Critério": finding.criterion_id,
                "Item do quadro": criterion_label(finding.criterion_id),
                "Elemento": finding.element_id,
                "Texto": finding.element_text,
                "Constatação": finding.title,
                "Tipo": finding.finding_type,
                "Gravidade proposta": finding.severity,
                "Confiança": finding.confidence,
                "Evidência": finding.evidence,
                "Explicação": finding.explanation,
                "Recomendação": finding.effective_recommendation,
                "Método": finding.method,
                "Decisão pericial": finding.expert_decision,
                "Gravidade final": finding.effective_severity,
                "Comentário pericial": finding.expert_comment,
                "Critérios relacionados": "; ".join(finding.related_criteria),
            }
        )
    return pd.DataFrame(rows)


def export_json(
    questionnaire: Questionnaire,
    metadata: QuestionnaireMetadata,
    findings: list[Finding],
    synthesis: Synthesis,
) -> bytes:
    payload = {
        "product": "SURV4IMPACT AI — Revisor de Questionários",
        "version": __version__,
        "generated_at": datetime.now().astimezone().isoformat(),
        "questionnaire": questionnaire.to_dict(),
        "metadata": metadata.to_dict(),
        "findings": [finding.to_dict() for finding in findings],
        "synthesis": synthesis.to_dict(),
    }
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


def export_xlsx(
    questionnaire: Questionnaire,
    metadata: QuestionnaireMetadata,
    findings: list[Finding],
    synthesis: Synthesis,
) -> bytes:
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        findings_dataframe(findings).to_excel(writer, sheet_name="Constatações", index=False)
        pd.DataFrame([metadata.to_dict()]).to_excel(writer, sheet_name="Contexto", index=False)
        pd.DataFrame([synthesis.to_dict()]).to_excel(writer, sheet_name="Síntese", index=False)
        pd.DataFrame([question.to_dict() for question in questionnaire.questions]).to_excel(
            writer, sheet_name="Perguntas", index=False
        )
        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns:
                letter = column[0].column_letter
                max_length = max(len(str(cell.value or "")) for cell in column)
                sheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 55)
    return buffer.getvalue()


def _set_cell_text(cell, text: str, *, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9)


def export_docx(
    questionnaire: Questionnaire,
    metadata: QuestionnaireMetadata,
    findings: list[Finding],
    synthesis: Synthesis,
) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    title = document.add_heading("SURV4IMPACT AI — Revisão metodológica", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(questionnaire.title)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.color.rgb = RGBColor(43, 91, 92)

    document.add_heading("Âmbito e ressalva", level=1)
    document.add_paragraph(
        "Relatório assistido por regras e, quando ativado, por IA. As constatações automáticas são "
        "propostas de apoio à decisão; a avaliação final permanece sob responsabilidade do avaliador."
    )

    document.add_heading("Contexto da aplicação", level=1)
    context_table = document.add_table(rows=0, cols=2)
    context_table.style = "Light Shading Accent 1"
    context_rows = [
        ("Modo", metadata.application_mode or "Não indicado"),
        ("Dispositivo", metadata.primary_device or "Não indicado"),
        ("Público-alvo", metadata.target_audience or "Não indicado"),
        ("Perfil", metadata.audience_profile or "Não indicado"),
        ("Necessidades de acessibilidade", metadata.accessibility_needs or "Não indicadas"),
    ]
    for label, value in context_rows:
        cells = context_table.add_row().cells
        _set_cell_text(cells[0], label, bold=True)
        _set_cell_text(cells[1], value)

    document.add_heading("Síntese do item D2.1.03", level=1)
    synthesis_paragraph = document.add_paragraph()
    synthesis_paragraph.add_run(synthesis.classification).bold = True
    synthesis_paragraph.add_run(f" — {synthesis.status}")
    document.add_paragraph(synthesis.rationale)
    document.add_paragraph(
        f"Validação: {synthesis.reviewed_findings}/{synthesis.total_findings} constatações revistas."
    )

    document.add_heading("Constatações", level=1)
    if not findings:
        document.add_paragraph("Não foram geradas constatações.")
    else:
        for index, finding in enumerate(findings, start=1):
            heading = document.add_paragraph()
            heading.add_run(
                f"{index}. {finding.criterion_id} — {finding.title}"
            ).bold = True
            document.add_paragraph(f"Elemento: {finding.element_id} — {finding.element_text}")
            document.add_paragraph(
                f"Gravidade: {finding.effective_severity} | Confiança proposta: {finding.confidence:.0%} "
                f"| Decisão: {finding.expert_decision}"
            )
            document.add_paragraph(f"Evidência: {finding.evidence}")
            document.add_paragraph(f"Leitura metodológica: {finding.explanation}")
            document.add_paragraph(f"Recomendação: {finding.effective_recommendation}")
            if finding.expert_comment:
                document.add_paragraph(f"Comentário do avaliador: {finding.expert_comment}")

    document.add_heading("Rastreabilidade", level=1)
    document.add_paragraph(
        "Cada constatação conserva o critério, a regra, o elemento analisado, a evidência, o método e a decisão pericial."
    )
    document.add_paragraph(
        "Gerado em " + datetime.now().astimezone().strftime("%d/%m/%Y %H:%M %Z")
    )

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
