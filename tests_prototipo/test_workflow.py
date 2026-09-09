"""Testes técnicos: respostas substitutas só existem nesta bateria, nunca na aplicação."""
from copy import deepcopy
from io import BytesIO
import json
import zipfile
import pytest
from docx import Document
from pydantic import ValidationError
from prototipo import api
from prototipo.engine import review_criterion, propose, validate_batch
from prototipo.models import Observation, Verification, ReviewBatch, CriterionProposal, Element
from prototipo.norma import load_norma, CRITERIOS
from prototipo.state import (new_session, transition, STAGES, replace_document, save_representation,
                            save_context, register_batch, answer, decide, completed,
                            withdraw_representation_confirmation)
from prototipo.extraction import extract, from_text
from prototipo.reporting import export_json, export_docx


def observation(rule="D2.4-R01", **kwargs):
    data = dict(criterio=rule[:4], regra_origem=rule, elemento="P1", evidencia_textual="Participou na oficina de escrita de 8 de setembro de 2026?",
        localizacao="P1", contexto_utilizado="contexto de teste técnico", observacao="Referente textual identificado na fixture técnica.",
        justificacao="Fixture de teste de contrato; não representa avaliação real.", consequencia_plausivel="Não avaliada neste teste técnico",
        alcance="P1", tipo_situacao="adequação documental", acao="nenhuma ação", estado="Concluída", informacao_adicional="",
        limites_inferencias_proibidas="Sem conclusão empírica ou validação metodológica", indispensavel_ao_nucleo=False,
        alteracao_necessaria="", proposta_redacao="", dependencias=["P1"])
    data.update(kwargs)
    return Observation(**data)


def verification(rule, deps=None, obs=None):
    return Verification(regra=rule, estado="Concluída", observacoes=obs or [], fundamento="Fixture técnica de cobertura", dependencias=deps or ["P1"], pendencia_metodologica="")


def fill(s):
    register_batch(s, ReviewBatch(verificacoes=[verification(r, obs=[observation(r)]) for r in load_norma()["regras"]]))


def proposal(c="D2.4", **kw):
    data = dict(criterio=c, estado="Concluída", categoria="Cumpre", evidencia_adequacao="Demonstração positiva na fixture",
        lacuna_motivo="", consequencia="", alcance="Instrumento sintético", atenuantes="", efeito_conjunto="",
        justificacao="Proposta substituta para testar decisão humana; não é apreciação real.", acao="", limites="Teste técnico")
    data.update(kw)
    return CriterionProposal(**data)


def test_norma_scope():
    n = load_norma()
    assert len(n["regras"]) == 24
    assert set(r[:4] for r in n["regras"]) == set(CRITERIOS)
    assert all(f"DM-0{i}" in n["comum"] for i in range(1, 6))
    assert n["versao"] == "v08" and len(n["sha256"]) == 64


def test_no_legacy_engine_imports():
    import ast
    from prototipo.norma import ROOT
    for p in (ROOT / "prototipo").glob("*.py"):
        imports = [n.module for n in ast.walk(ast.parse(p.read_text(encoding="utf-8"))) if isinstance(n, ast.ImportFrom)]
        assert not set(imports) & {"surv4impact.rules", "surv4impact.synthesis", "surv4impact.knowledge", "surv4impact.ai"}


def test_transition_validated_before_write(session):
    old = session["stage"]
    with pytest.raises(ValueError):
        transition(session, "Confirmar representação")
    assert session["stage"] == old
    for target in STAGES:
        transition(session, target)
        assert session["stage"] == target


def test_gate_before_confirmation():
    s = new_session()
    with pytest.raises(ValueError):
        transition(s, "structure")
    instrument, elements = from_text("P1. Pergunta")
    replace_document(s, instrument, elements)
    with pytest.raises(ValueError):
        transition(s, "review")
    with pytest.raises(ValueError):
        review_criterion(s, "D2.4", "", "fixture")


def test_replace_and_reset_no_residual(session):
    fill(session)
    session["decisoes"]["D2.4"] = {"categoria": "Cumpre"}
    previous_id = session["id"]
    doc, elems = from_text("Novo instrumento")
    replace_document(session, doc, elems)
    assert session["id"] != previous_id
    assert not session["verificacoes"] and not session["decisoes"] and not session["contexto"]
    assert not session["confirmada"]


def test_unknown_format_preserved_after_parse_and_confirmation():
    d = Document()
    d.add_paragraph("1. Selecione a situação aplicável.")
    buffer = BytesIO(); d.save(buffer)
    instrument, elems = extract(buffer.getvalue(), "unknown.docx")
    assert all(e["formato"] == "não confirmado" for e in elems)
    s = new_session(); replace_document(s, instrument, elems)
    save_representation(s, elems, True)
    assert all(e["formato"] == "não confirmado" for e in s["representacao"])


def test_matrices_and_shared_instructions_are_preserved(session):
    elems = deepcopy(session["representacao"])
    elems[2].update(formato="matriz", linhas=["clareza", "utilidade"], colunas=["Baixa", "Alta"], subitens=["P2a", "P2b"], opcoes=["Baixa", "Alta"])
    elems[0].update(instrucoes="Nas últimas quatro semanas", instrucao_aplica_a=["P1", "P2", "P3"])
    save_representation(session, elems, True)
    restored = json.loads(export_json(session))
    assert restored["representacao"][2]["linhas"] == ["clareza", "utilidade"]
    assert restored["representacao"][0]["instrucao_aplica_a"] == ["P1", "P2", "P3"]
    assert restored["representacao"][1]["condicoes_destinos"] == "Sim → P2; Não / Prefiro não responder → fim"


def test_raw_docx_table_locations():
    d = Document(); d.add_paragraph("Instrução comum")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Dimensão"; t.cell(0, 1).text = "Sim"
    t.cell(1, 0).text = "Clareza"; t.cell(1, 1).text = "[ ]"
    b = BytesIO(); d.save(b)
    instrument, _ = extract(b.getvalue(), "matrix.docx")
    assert instrument["texto_bruto"][1] == {"localizacao": "DOCX bloco 2", "texto": "Dimensão | Sim\nClareza | [ ]"}


def test_correction_invalidates_only_dependencies(session):
    register_batch(session, ReviewBatch(verificacoes=[verification("D2.3-R05", ["P1"]), verification("D2.4-R01", ["P2"])]))
    session["propostas"]["D2.4"] = proposal().model_dump()
    elems = deepcopy(session["representacao"]); elems[1]["texto"] += " Texto corrigido."
    version = session["instrumento"]["versao"]
    save_representation(session, elems, True)
    assert "D2.3-R05" not in session["verificacoes"] and "D2.4-R01" in session["verificacoes"]
    assert "D2.4" in session["propostas"]
    assert session["instrumento"]["versao"] == version
    assert session["historico"][-1]["tipo"] == "correção da representação"


def test_unconfirmed_correction_preserves_independent_results(session):
    register_batch(session, ReviewBatch(verificacoes=[
        verification("D2.3-R05", ["P1"]),
        verification("D2.4-R01", ["P2"]),
    ]))
    session["propostas"]["D2.4"] = proposal().model_dump()
    elements = deepcopy(session["representacao"])
    elements[1]["texto"] += " Texto corrigido antes de nova confirmação."

    save_representation(session, elements, False)

    assert not session["confirmada"]
    assert not session["confirmacao"]["explicita"]
    assert "D2.3-R05" not in session["verificacoes"]
    assert "D2.4-R01" in session["verificacoes"]
    assert "D2.4" in session["propostas"]


def test_withdrawing_confirmation_does_not_change_reading_or_results(session):
    register_batch(session, ReviewBatch(verificacoes=[verification("D2.4-R01", ["P2"])]))
    before_reading = deepcopy(session["representacao"])
    before_results = deepcopy(session["verificacoes"])

    withdraw_representation_confirmation(session)

    assert not session["confirmada"]
    assert not session["confirmacao"]["explicita"]
    assert session["representacao"] == before_reading
    assert session["verificacoes"] == before_results


def test_context_visibility_and_invalidation(session):
    register_batch(session, ReviewBatch(verificacoes=[verification("D2.4-R05", ["contexto:período"]), verification("D2.3-R05", ["P1"])]))
    context = session["contexto"] | {"período": {"valor": "últimos seis meses", "visibilidade": "contexto interno fornecido à aplicação"}}
    save_context(session, context, [])
    assert "D2.4-R05" not in session["verificacoes"] and "D2.3-R05" in session["verificacoes"]
    assert "últimos seis meses" not in session["representacao"][1]["texto"]


def test_shared_instruction_change_invalidates_inheriting_items(session):
    register_batch(session, ReviewBatch(verificacoes=[verification("D2.4-R05", ["P2"])]))
    assert "INTRO" in session["verificacoes"]["D2.4-R05"]["dependencias"]
    elements = deepcopy(session["representacao"])
    elements[0]["texto"] += " Instrução corrigida do original."
    save_representation(session, elements, True)
    assert "D2.4-R05" not in session["verificacoes"]


def test_added_element_invalidates_coverage(session):
    fill(session)
    elements = deepcopy(session["representacao"])
    elements.append(Element(id="P4", texto="Nova pergunta", localizacao="P4").model_dump())
    save_representation(session, elements, True)
    assert not session["verificacoes"]


def test_withdrawn_action_cannot_remain_current_after_correction(session):
    obs = observation(tipo_situacao="problema demonstrado", acao="correção necessária", alteracao_necessaria="Fixture de alteração")
    register_batch(session, ReviewBatch(verificacoes=[verification("D2.4-R01", obs=[obs])]))
    action = next(iter(session["acoes"].values()))
    action["aceite"] = True
    elements = deepcopy(session["representacao"]); elements[1]["texto"] += " correção de leitura"
    save_representation(session, elements, True)
    assert action["estado"] == "retirada para reapreciação"
    assert action["incorporada"] is False


def test_absence_of_alerts_cannot_propose_compliance(session):
    register_batch(session, ReviewBatch(verificacoes=[verification(r) for r in load_norma()["regras"]]))
    with pytest.raises(api.TechnicalError, match="positiva"):
        propose(session, "D2.4", "fixture", "fixture", lambda *a: proposal())


def test_explicit_free_only_d26_exemption(session):
    elements = [session["representacao"][0], Element(id="P1", texto="Comentário livre", localizacao="P1", formato="livre", restricoes="nenhuma").model_dump()]
    save_representation(session, elements, True)
    fill(session)
    result = propose(session, "D2.6", "fixture", "fixture", lambda *a: proposal("D2.6", categoria="Não Aplicável"))
    assert result.categoria == "Não Aplicável" and not session["decisoes"]


def test_no_automatic_repetition_of_closed_request(session):
    obs = observation(tipo_situacao="documentação/contexto não disponibilizado", acao="pedido de esclarecimento", informacao_adicional="Qual o destino?", estado="A aguardar esclarecimento")
    batch = ReviewBatch(verificacoes=[verification("D2.4-R01", obs=[obs])])
    register_batch(session, batch)
    answer(session, session["pedidos"][0]["id"], "Não disponível", "documento não disponível", unavailable=True)
    with pytest.raises(api.TechnicalError, match="já fechado"):
        validate_batch(session, batch, ["D2.4-R01"])


@pytest.mark.parametrize("unavailable,close", [(False, False), (True, False), (False, True)])
def test_request_answer_and_explicit_closure(session, unavailable, close):
    o = observation(tipo_situacao="esclarecimento necessário", acao="pedido de esclarecimento", estado="A aguardar esclarecimento", informacao_adicional="Qual a unidade de resposta?", indispensavel_ao_nucleo=True)
    register_batch(session, ReviewBatch(verificacoes=[verification("D2.4-R01", obs=[o]), verification("D2.3-R05")]))
    p = session["pedidos"][0]
    answer(session, p["id"], "Não disponível" if unavailable or close else "O participante", "contexto interno fornecido à aplicação", unavailable, close)
    assert p["estado"] == ("fechado com limites" if unavailable or close else "respondido — por reapreciar")
    assert "D2.4-R01" not in session["verificacoes"] and "D2.3-R05" in session["verificacoes"]
    assert not session["propostas"] and not session["decisoes"]
    assert not session["acoes"][p["id"]]["incorporada"]


def test_silence_does_not_close(session):
    o = observation(informacao_adicional="Qual o destino?", estado="A aguardar esclarecimento", acao="pedido de esclarecimento", tipo_situacao="esclarecimento necessário")
    register_batch(session, ReviewBatch(verificacoes=[verification("D2.4-R01", obs=[o])]))
    with pytest.raises(ValueError):
        answer(session, session["pedidos"][0]["id"], "", "contexto interno fornecido à aplicação", close=True)
    assert session["pedidos"][0]["estado"] == "pendente"


def test_api_missing_and_failure_are_technical(session, monkeypatch):
    status = api.smoke_test("", "configured-model")
    assert status["sucesso"] is False and status["latencia_s"] is None
    def failure(*args, **kw):
        raise RuntimeError("secret-example-do-not-display")
    monkeypatch.setattr(api, "OpenAI", failure)
    status = api.smoke_test("secret-example-do-not-display", "configured-model")
    assert "secret-example" not in json.dumps(status)
    with pytest.raises(api.TechnicalError, match="RuntimeError"):
        api.structured("secret", "configured-model", "", {}, ReviewBatch)
    assert not session["verificacoes"] and not session["propostas"]


def test_reject_unverifiable_evidence(session):
    o = observation(evidencia_textual="conteúdo inventado")
    with pytest.raises(api.TechnicalError, match="Evidência"):
        validate_batch(session, ReviewBatch(verificacoes=[verification("D2.4-R01", obs=[o])]), ["D2.4-R01"])


def test_positive_observation_not_discarded(session):
    batch = ReviewBatch(verificacoes=[verification("D2.4-R01", obs=[observation()])])
    validate_batch(session, batch, ["D2.4-R01"])
    register_batch(session, batch)
    assert session["verificacoes"]["D2.4-R01"]["observacoes"][0]["tipo_situacao"] == "adequação documental"


def test_normal_prompt_excludes_test_documents(session):
    def fake(key, model, instructions, data, schema):
        assert "casos_teste_metodologicos" not in instructions
        assert "CT-01A" not in instructions
        assert set(data["regras_a_verificar"]) == {r for r in load_norma()["regras"] if r.startswith("D2.4")}
        return ReviewBatch(verificacoes=[verification(r) for r in data["regras_a_verificar"]])
    review_criterion(session, "D2.4", "fixture", "fixture", fake)
    assert set(session["verificacoes"]) == {r for r in load_norma()["regras"] if r.startswith("D2.4")}


def test_classification_requires_human(session):
    fill(session)
    propose(session, "D2.4", "fixture", "fixture", lambda *a: proposal())
    assert not session["decisoes"] and not completed(session)
    with pytest.raises(ValueError):
        decide(session, "D2.4", "Cumpre", "Texto", "Avaliador de teste", False)
    decide(session, "D2.4", "Cumpre", "Aceitação expressa da fixture técnica", "Avaliador de teste", True)
    assert session["decisoes"]["D2.4"]["expressa"]


def test_local_scope_cannot_classify(session):
    fill(session); session["ambito"] = "recorte localizado"
    with pytest.raises(ValueError, match="recorte"):
        propose(session, "D2.4", "fixture", "fixture", lambda *a: proposal())


def test_pending_essential_blocks_category(session):
    fill(session)
    o = observation(indispensavel_ao_nucleo=True, estado="A aguardar validação da extração", tipo_situacao="dúvida de extração ou representação")
    register_batch(session, ReviewBatch(verificacoes=[verification("D2.4-R01", obs=[o])]))
    with pytest.raises(api.TechnicalError, match="pendente"):
        propose(session, "D2.4", "fixture", "fixture", lambda *a: proposal())


def test_unknown_cannot_be_not_applicable(session):
    fill(session)
    with pytest.raises(api.TechnicalError, match="inexistência"):
        propose(session, "D2.6", "fixture", "fixture", lambda *a: proposal("D2.6", categoria="Não Aplicável"))


def test_state_is_not_category():
    with pytest.raises(ValidationError):
        proposal(categoria="A aguardar esclarecimento")


def test_export_partial_complete_and_no_global(session):
    partial = json.loads(export_json(session))
    assert partial["tipo_relatorio"] == "Parcial"
    fill(session)
    for c in CRITERIOS:
        propose(session, c, "fixture", "fixture", lambda *a, c=c: proposal(c))
        decide(session, c, "Cumpre", "Decisão expressa para fixture técnica", "Avaliador de teste", True)
    assert completed(session)
    full = json.loads(export_json(session))
    assert full["tipo_relatorio"] == "Concluído" and full["criterios_por_concluir"] == []
    assert "classificacao_global" not in full
    docx = export_docx(session)
    assert zipfile.is_zipfile(BytesIO(docx))
    document = Document(BytesIO(docx))
    assert "SURV4IMPACT AI" in "\n".join(p.text for p in document.paragraphs)
    assert len(document.inline_shapes) == 3
