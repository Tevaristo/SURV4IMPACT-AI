"""Conteúdo/proveniência do exemplo; não simula validação semântica por uma API."""
import hashlib
from io import BytesIO
import json
import re
from pathlib import Path
from unittest.mock import Mock

import pytest
from openpyxl import load_workbook

from prototipo.demo import (BASE_METADATA, BASE_RESOURCE, RESOURCE, METADATA,
                            EXAMPLE_ID, EXAMPLE_TITLE, example_content_sha256,
                            load_example, load_demo, example_provenance)
from prototipo.engine import payload, propose, review_criterion
from prototipo.models import Element, ReviewBatch, Verification
from prototipo.norma import NORMA, ROOT, CRITERIOS, load_norma
from prototipo.state import new_session, save_representation
from prototipo.table_export import READING_COLUMNS, export_reading_csv, export_reading_xlsx
from .cases import APPROVED, DIRECT_RULES, SOURCE


def case_section(code, letter):
    source = SOURCE.read_text(encoding="utf-8-sig")
    body = re.search(rf"(?ms)^## {code} .*?(?=^## |\Z)", source).group()
    return re.search(rf"(?ms)^### {letter}\..*?(?=^### |\Z)", body).group()


def test_demo_sources_and_approved_status():
    metadata = example_provenance()
    assert metadata["conteudo_sha256"] == example_content_sha256()
    assert metadata["casos_aprovados_utilizados"] == ["CT-03", "CT-04A", "CT-08A", "CT-12A", "CT-13"]
    assert set(metadata["casos_aprovados_utilizados"]) <= set(APPROVED)
    assert set(metadata["fontes"]) == {NORMA.name, SOURCE.name,
        "validacao_metodologica_casos_v02.md", "inventario_documentos_teste_v11.md",
        "validacao_metodologica_ct12a_v03.md"}
    for name, origin in metadata["fontes"].items():
        path = NORMA if name == NORMA.name else SOURCE.parent / name
        assert origin["sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
    validation = (SOURCE.parent / "validacao_metodologica_casos_v02.md").read_text(encoding="utf-8-sig")
    for group in metadata["grupos"]:
        code = group["caso"]
        if code == "CT-12A":
            decision_path = SOURCE.parent / group["validacao_origem"]["fonte"]
            assert group["validacao_origem"]["estatuto"] == "Decisão metodológica posterior da responsável"
            assert group["validacao_origem"]["historico_preservado"] == "validacao_metodologica_casos_v02.md"
            assert group["validacao_origem"]["sha256_ficha_validacao"] == hashlib.sha256(decision_path.read_bytes()).hexdigest()
            assert group["proveniencia_entrada_expectativas"]["sha256_documento"] == hashlib.sha256(decision_path.read_bytes()).hexdigest()
            assert group["regras"] == DIRECT_RULES[code]
            continue
        decision = re.search(rf"(?ms)^## {code} .*?(?=^## |\Z)", validation).group()
        assert "**Estado da validação V11:** Aprovado." in decision
        assert group["validacao_origem"]["estatuto"] == "Aprovado"
        assert group["validacao_origem"]["sha256_ficha_validacao"] == hashlib.sha256(decision.encode()).hexdigest()
        assert group["regras"] == DIRECT_RULES[code]
        for letter, digest in group["proveniencia_entrada_expectativas"]["sha256_secoes"].items():
            assert digest == hashlib.sha256(case_section(code, letter).encode()).hexdigest()
    assert metadata["criterios_abrangidos"] == list(CRITERIOS)
    assert set(metadata["regras_abrangidas"]) == {r for g in metadata["grupos"] for r in g["regras"]}
    assert set(metadata["regras_abrangidas"]) <= set(load_norma()["regras"])


def test_internal_demo_loads_from_another_directory_without_case_files(monkeypatch, tmp_path):
    monkeypatch.chdir(tmp_path)
    original = Path.read_text
    reads = []
    def guarded(path, *args, **kwargs):
        assert "base de conhecimento" not in str(path)
        reads.append(path)
        return original(path, *args, **kwargs)
    with monkeypatch.context() as patch:
        patch.setattr(Path, "read_text", guarded)
        example = load_example()
        metadata = example_provenance()
        assert set(reads) == {BASE_RESOURCE, RESOURCE, BASE_METADATA, METADATA}
    assert example["id"] == EXAMPLE_ID and example["titulo"] == EXAMPLE_TITLE
    assert metadata["id"] == example["id"]
    assert "C:\\" not in json.dumps(example) and "base de conhecimento" not in json.dumps(example)


def test_all_demo_elements_have_traceability_and_visible_source_text():
    example, metadata = load_example(), example_provenance()
    elements = {e["id"]: Element.model_validate(e) for e in example["elementos"]}
    assert len(elements) == len(example["elementos"])
    assert set(elements) == {m["elemento"] for m in metadata["correspondencia"]}
    for entry in metadata["correspondencia"]:
        assert entry["caso"] in APPROVED and entry["localizacao_original"]
        assert entry["regras"] == DIRECT_RULES[entry["caso"]]
    for element in elements.values():
        assert element.texto in example["texto"]
        assert all(v in example["texto"] for field in ("opcoes", "subitens", "linhas", "colunas") for v in getattr(element, field))
        assert set(element.instrucao_aplica_a) <= set(elements)
    assert "Garantimos a confidencialidade das suas respostas" in elements["INTRO"].texto
    assert elements["P1"].texto == "Qual a entidade a que pertence?"
    assert elements["P8"].texto == "Comentários e sugestões"
    assert len(elements["SAT"].instrucao_aplica_a) == 8
    for element_id in elements["SAT"].instrucao_aplica_a:
        item = elements[element_id]
        assert item.formato == "escolha única"
        assert item.opcoes == ["1", "2", "3", "4", "5"]
        assert item.regra_selecao == "Uma única resposta"
        assert item.numero_minimo == item.numero_maximo == 1
        assert "1 pouco satisfeito" in item.instrucoes and "5 muito satisfeito" in item.instrucoes
    assert "relevância" in elements["P3"].texto and "0 a 100" in elements["P3"].texto
    assert "sucesso" in elements["P4"].texto and "0 a 100" in elements["P4"].texto
    assert elements["P5"].texto in case_section("CT-08A", "B")
    assert elements["P5"].linhas == ["Expectativa sobre a execução dos valores de investimento contratados"]
    assert len(elements["P5"].colunas) == 4
    assert all(v in case_section("CT-08A", "B") for v in elements["P5"].linhas + elements["P5"].colunas)
    assert elements["P5"].regra_selecao == "Outra condição indicada no questionário"
    assert elements["P5"].outra_condicao_resposta == "Uma resposta por linha"
    assert elements["P5"].numero_minimo is None and elements["P5"].numero_maximo is None
    assert "P6.2" in elements["P6.3"].condicoes_destinos
    assert elements["P6.2"].opcoes == ["Sim", "Não"]
    assert "Opções:\n- Sim\n- Não\n\nP6.3." in example["texto"]
    assert len(elements["P6.4"].opcoes) == 5
    assert all(v in case_section("CT-13", "B") for v in elements["P6.4"].opcoes)
    assert "menor dimensão financeira; redução no valor de investimento em %" in elements["P6.4"].opcoes[1]
    assert elements["P6.4"].regra_selecao == "Uma única resposta"
    assert elements["P6.4"].numero_minimo == elements["P6.4"].numero_maximo == 1
    unknown = elements["P7"]
    assert unknown.formato == "não confirmado" and len(unknown.opcoes) == 4
    assert unknown.regra_selecao == "Não foi possível confirmar"
    assert unknown.numero_minimo is None and unknown.numero_maximo is None
    assert not unknown.subitens and not unknown.linhas and not unknown.colunas
    normalized_source = " ".join(case_section("CT-12A", "B").split())
    assert all(" ".join(v.split()) in normalized_source for v in unknown.opcoes)
    assert all(elements[element_id].regra_selecao == "Não aplicável" for element_id in ("INTRO", "SAT", "APOIO", "FIM"))
    for element_id in ("P3", "P4"):
        item = elements[element_id]
        assert "0 a 100" in item.texto
        assert item.regra_selecao == "Não aplicável"
        assert item.numero_minimo is None and item.numero_maximo is None


def test_expected_behaviours_are_localized_propositions_not_categories():
    metadata = example_provenance()
    expected = {g["caso"]: g["esperado"] for g in metadata["grupos"]}
    assert metadata["classificacao_integral_validada"] is False
    assert metadata["alcance"] == "recorte localizado"
    for value in expected.values():
        assert value["proposicoes"] and value["inferencias_proibidas"]
        assert "classificar_criterio" in value["inferencias_proibidas"]
        assert "categoria" not in value
    for code in ("CT-04A", "CT-08A", "CT-13"):
        assert expected[code]["acao"] == "nenhuma ação"
        assert expected[code]["estados"] == ["Concluída"]
        assert expected[code]["conclusao_positiva_localizada"]
    assert expected["CT-03"]["conclusao_positiva_localizada"]
    assert not expected["CT-03"]["apreciacao_integral_regra_concluida"]
    assert expected["CT-03"]["acao"] == "pedido de esclarecimento"
    assert "uso_divulgacao_respostas_comentarios_insuficientemente_documentados" in expected["CT-03"]["proposicoes"]
    assert expected["CT-12A"]["tipo_situacao"] == "problema demonstrado"
    assert "quatro_alternativas_presentes" in expected["CT-12A"]["proposicoes"]
    assert "regra_selecao_ausente_diretamente_observavel" in expected["CT-12A"]["proposicoes"]
    assert "pedir_ou_criar_escala_sim_nao" in expected["CT-12A"]["inferencias_proibidas"]
    assert expected["CT-12A"]["classificacao_integral_determinada"] is False
    assert expected["CT-12A"]["evidencia"] == "Existem quatro alternativas e não existe uma instrução que indique quantas podem ou devem ser selecionadas."
    assert expected["CT-12A"]["problema"] == "A regra de seleção não está definida."
    assert "não são mutuamente exclusivas" in expected["CT-12A"]["justificacao"]
    assert "regras de resposta distintas" in expected["CT-12A"]["consequencia"]
    assert expected["CT-12A"]["acao_dirigida"] == "Esclarecer se deve ser selecionada uma única alternativa ou todas as alternativas aplicáveis."
    assert expected["CT-12A"]["propostas_condicionais"]["escolha_multipla"] == "Selecione todas as opções aplicáveis."
    assert expected["CT-12A"]["propostas_condicionais"]["escolha_unica"] == "Selecione a opção que melhor descreve a situação mais frequente."
    assert "exige validação metodológica" in expected["CT-12A"]["propostas_condicionais"]["limite_escolha_unica"]
    assert expected["CT-12A"]["acao"] == "pedido de esclarecimento"
    assert expected["CT-12A"]["estados"] == ["A aguardar esclarecimento"]


def test_demo_table_exports_preserve_multiline_options_and_structured_limits():
    elements = [Element.model_validate(value).model_dump() for value in load_example()["elementos"]]
    index = {key: position for position, (key, _) in enumerate(READING_COLUMNS, 1)}
    row = {element["id"]: position for position, element in enumerate(elements, 2)}
    workbook = load_workbook(BytesIO(export_reading_xlsx(elements)), data_only=False)
    sheet = workbook["Leitura do questionário"]
    assert sheet.cell(row["P2a"], index["opcoes"]).value == "1\n2\n3\n4\n5"
    assert sheet.cell(row["P6.2"], index["opcoes"]).value == "Sim\nNão"
    assert sheet.cell(row["P6.4"], index["opcoes"]).value.splitlines()[1] == "Sim, mas com alterações: menor dimensão financeira; redução no valor de investimento em %"
    assert sheet.cell(row["P7"], index["opcoes"]).value == "\n".join(next(element for element in elements if element["id"] == "P7")["opcoes"])
    assert sheet.cell(row["P7"], index["numero_minimo"]).value is None
    csv_text = export_reading_csv(elements).decode("utf-8-sig")
    assert '"1\n2\n3\n4\n5"' in csv_text
    assert '"Sim\nNão"' in csv_text


def test_demo_resets_state_without_results_and_cannot_classify():
    s = new_session()
    old_id = s["id"]
    s["propostas"] = {"D2.4": {"anterior": True}}
    load_demo(s)
    assert s["id"] != old_id and s["stage"] == "structure"
    assert s["instrumento"]["exemplo_id"] == EXAMPLE_ID
    assert s["instrumento"]["texto_bruto"][0]["texto"] == load_example()["texto"]
    assert not s["confirmada"]
    assert all(c["visibilidade"] == "contexto interno fornecido à aplicação" for c in s["contexto"].values())
    for field in ("verificacoes", "propostas", "decisoes", "pedidos", "acoes", "historico"):
        assert not s[field]
    save_representation(s, s["representacao"], True)
    # Cobertura fictícia apenas para testar o bloqueio preexistente do alcance.
    s["verificacoes"] = {r: {} for r in load_norma()["regras"]}
    caller = Mock(side_effect=AssertionError("Não pode chamar a API para classificar o recorte"))
    for criterion in CRITERIOS:
        with pytest.raises(ValueError, match="recorte localizado"):
            propose(s, criterion, "", "", caller)
    caller.assert_not_called()
    assert not s["propostas"] and not s["decisoes"]


def test_demo_expectations_never_enter_normal_analysis_payload():
    s = new_session()
    load_demo(s)
    save_representation(s, s["representacao"], True)
    data = json.dumps(payload(s), ensure_ascii=False)
    for group in example_provenance()["grupos"]:
        assert group["caso"] not in data
        assert all(p not in data for p in group["esperado"]["proposicoes"])
    def inspect_only(key, model, instructions, sent, schema):
        assert "grupos" not in sent and "correspondencia" not in sent
        assert "esperado" not in sent
        assert sent["ambito"] == "recorte localizado"
        assert not sent["pedidos"] and not sent["acoes"]
        # Fixture de contrato, sem observações nem conclusão metodológica.
        return ReviewBatch(verificacoes=[Verification(regra=r, estado="Por concluir", observacoes=[],
            fundamento="Teste técnico do payload", dependencias=["P7"], pendencia_metodologica="")
            for r in sent["regras_a_verificar"]])
    review_criterion(s, "D2.6", "", "", inspect_only)
    assert not s["propostas"]
