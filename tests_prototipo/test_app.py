from copy import deepcopy
from hashlib import sha256
from io import BytesIO

import pytest
from docx import Document
from streamlit.testing.v1 import AppTest
from prototipo.demo import TEXT, EXAMPLE_DESCRIPTION, EXAMPLE_LIMIT
from prototipo.engine import payload
from prototipo.norma import ROOT
from prototipo.reporting import export_json, sections
from prototipo.state import save_representation
from .test_brand import assert_native_header


def app():
    return AppTest.from_file(str(ROOT / "app_prototipo.py"), default_timeout=30).run()


def click_label(a, label):
    next(b for b in a.button if b.label == label).click().run()
    assert not a.exception


def test_startup_and_institutional_identity():
    a = app()
    assert not a.exception
    assert_native_header(a)


def test_navigation_full_partial_report_and_new_analysis():
    a = app()
    a.button(key="nav_review").click().run()
    assert a.session_state["prototype"]["stage"] == "intake"
    a.button(key="demo").click().run()
    assert a.session_state["prototype"]["stage"] == "structure"
    assert a.button(key="nav_structure").label == "Rever a leitura do questionário"
    assert a.header[0].value == "Rever a leitura do questionário"
    click_label(a, "Guardar leitura e confirmação")
    assert not a.session_state["prototype"]["confirmada"]
    assert next(b for b in a.button if b.label == "Continuar para apreciação").disabled
    a.button(key="nav_review").click().run()
    assert a.session_state["prototype"]["stage"] == "structure"
    assert any(w.value == "Confirme que a leitura do questionário está correta antes de iniciar a análise." for w in a.warning)
    next(c for c in a.checkbox if c.label.startswith("Confirmo que a leitura está correta")).check()
    click_label(a, "Guardar leitura e confirmação")
    assert a.session_state["prototype"]["confirmada"]
    for stage in ("intake", "structure", "review", "clarifications", "decisions", "report"):
        a.button(key="nav_" + stage).click().run()
        assert not a.exception
        assert a.session_state["prototype"]["stage"] == stage
        assert_native_header(a)
    assert any("Relatório parcial" in t.value for t in a.markdown)
    a.button(key="reset").click().run()
    assert a.session_state["prototype"]["stage"] == "intake"
    assert not a.session_state["prototype"]["instrumento"]
    assert not a.exception


def test_document_substitution_by_text_has_fresh_widgets():
    a = app()
    a.button(key="demo").click().run()
    old_id = a.session_state["prototype"]["id"]
    a.button(key="nav_intake").click().run()
    next(t for t in a.text_area if t.label == "Representação textual fiel (alternativa ao ficheiro)").input("P1. Novo questionário sem opções confirmadas.")
    click_label(a, "Extrair / substituir instrumento")
    s = a.session_state["prototype"]
    assert s["id"] != old_id and not s["confirmada"]
    assert "Novo questionário" in s["representacao"][0]["texto"]
    assert s["representacao"][0]["formato"] == "não confirmado"


def test_full_reading_replacement_is_explicit_and_requires_new_confirmation():
    from .test_workflow import verification
    a = app()
    a.button(key="demo").click().run()
    s = a.session_state["prototype"]
    original = deepcopy(s["representacao"])
    instrument = deepcopy(s["instrumento"])
    panel = next(e for e in a.expander if e.label == "Substituir toda a leitura automática (opcional)")
    assert not panel.proto.expanded
    assert len(panel.text_area) == 1
    assert panel.text_area[0].label == "Conteúdo completo do questionário"
    assert panel.selectbox[0].label == "Localização no documento original (opcional)"
    assert panel.selectbox[0].value is None
    assert panel.selectbox[0].proto.accept_new_options
    assert panel.selectbox[0].proto.placeholder == "Ex.: páginas 4–8, Bloco B — Satisfação"
    assert [b.label for b in panel.button] == ["Substituir a leitura automática"]
    assert len(panel.warning) == 1
    replacement = "Instruções: escolha uma opção.\nP1. Participou na oficina?\nSim\nNão\nFim."
    panel.text_area[0].input(replacement)
    panel.selectbox[0].set_value("Questionário completo")
    a.run()
    assert s["representacao"] == original
    next(c for c in a.checkbox if c.label.startswith("Confirmo que a leitura")).check()
    click_label(a, "Guardar leitura e confirmação")
    assert s["representacao"] == original and s["confirmada"]

    # Resultados técnicos com dependências explícitas, sem apreciação simulada.
    independent = verification("D2.4-R06", ["contexto:finalidade"]).model_dump()
    s["verificacoes"] = {"D2.4-R01": verification("D2.4-R01", ["P1"]).model_dump(),
                         "D2.4-R06": deepcopy(independent)}
    click_label(a, "Substituir a leitura automática")
    assert len(s["representacao"]) == 1
    assert s["representacao"][0]["texto"] == replacement
    assert s["representacao"][0]["localizacao"] == "Conteúdo introduzido pelo utilizador"
    assert s["representacao"][0]["formato"] == "não confirmado"
    assert s["instrumento"] == instrument
    assert s["verificacoes"] == {"D2.4-R06": independent}
    assert not s["confirmada"] and not s["confirmacao"]["explicita"]
    assert s["rastreabilidade_substituicao"]["localizacao_documento_original"] == "Questionário completo"
    assert "Questionário completo" in export_json(s).decode("utf-8")
    assert any("Questionário completo" in line for _, lines in sections(s) for line in lines)
    assert payload(s)["representacao"][0]["localizacao"] == "Conteúdo introduzido pelo utilizador"
    assert "rastreabilidade_substituicao" not in payload(s)
    confirmation = next(c for c in a.checkbox if c.label.startswith("Confirmo que a leitura"))
    assert not confirmation.value
    assert next(b for b in a.button if b.label == "Continuar para apreciação").disabled
    a.button(key="nav_review").click().run()
    assert s["stage"] == "structure"
    next(c for c in a.checkbox if c.label.startswith("Confirmo que a leitura")).check()
    click_label(a, "Guardar leitura e confirmação")
    assert s["confirmada"]

    # Mesmo texto: a nova substituição não pode herdar a confirmação anterior.
    revision = s["revisao"]
    next(t for t in a.text_area if t.label == "Conteúdo completo do questionário").input(replacement)
    click_label(a, "Substituir a leitura automática")
    assert s["revisao"] == revision
    assert not s["confirmada"]
    assert not next(c for c in a.checkbox if c.label.startswith("Confirmo que a leitura")).value
    assert s["verificacoes"] == {"D2.4-R06": independent}


def test_empty_full_reading_replacement_preserves_current_content():
    a = app()
    a.button(key="demo").click().run()
    next(c for c in a.checkbox if c.label.startswith("Confirmo que a leitura")).check()
    click_label(a, "Guardar leitura e confirmação")
    before = deepcopy(a.session_state["prototype"])
    next(t for t in a.text_area if t.label == "Conteúdo completo do questionário").input("   ")
    click_label(a, "Substituir a leitura automática")
    assert a.error
    assert a.session_state["prototype"] == before


@pytest.mark.parametrize("location, expected", [
    (None, "Localização não indicada pelo utilizador"),
    ("Questionário completo", "Questionário completo"),
])
def test_optional_replacement_location_is_traceability_only(location, expected):
    a = app()
    a.button(key="demo").click().run()
    s = a.session_state["prototype"]
    replacement = "Conteúdo integral para teste técnico."
    next(t for t in a.text_area if t.label == "Conteúdo completo do questionário").input(replacement)
    location_field = next(x for x in a.selectbox if x.label == "Localização no documento original (opcional)")
    if location is not None:
        location_field.set_value(location)
    click_label(a, "Substituir a leitura automática")
    assert s["rastreabilidade_substituicao"]["localizacao_documento_original"] == expected
    assert expected in export_json(s).decode("utf-8")
    assert any(expected in line for _, lines in sections(s) for line in lines)
    assert "rastreabilidade_substituicao" not in payload(s)
    assert payload(s)["representacao"][0]["localizacao"] == "Conteúdo introduzido pelo utilizador"
    assert not s["confirmada"]


def test_pdf_upload_replacement_raw_pages_and_title(prototype_pdfs):
    first, second = prototype_pdfs
    a = app()
    a.file_uploader[0].set_value(("primeiro.pdf", first, "application/pdf"))
    a.run(); click_label(a, "Extrair / substituir instrumento")
    s = a.session_state["prototype"]
    first_id = s["id"]
    assert s["instrumento"]["designacao"] == "Questionário PDF de teste"
    assert [b["localizacao"] for b in s["instrumento"]["texto_bruto"]] == ["PDF p. 1", "PDF p. 2"]
    assert "Clareza" in s["instrumento"]["texto_bruto"][1]["texto"]
    assert all(e["formato"] == "não confirmado" for e in s["representacao"])
    a.file_uploader[0].set_value(("segundo.pdf", second, "application/pdf"))
    a.run(); click_label(a, "Extrair / substituir instrumento")
    s = a.session_state["prototype"]
    assert s["id"] != first_id
    assert s["instrumento"]["designacao"] == "Segundo questionário PDF"
    assert not s["confirmada"] and not s["verificacoes"] and not s["contexto"]


@pytest.mark.parametrize("extension", ["pdf", "docx"])
def test_example_and_upload_replace_each_other(extension, prototype_pdfs):
    if extension == "pdf":
        content = prototype_pdfs[0]
        mime = "application/pdf"
    else:
        doc = Document()
        doc.add_heading("Questionário do utilizador", 0)
        doc.add_paragraph("P1. Participou na oficina?")
        doc.add_paragraph("Sim")
        doc.add_paragraph("Não")
        buffer = BytesIO()
        doc.save(buffer)
        content = buffer.getvalue()
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    upload = ("questionario." + extension, content, mime)
    a = app()
    assert a.button(key="demo").label == "Usar questionário de exemplo"
    uploader = a.file_uploader[0]
    assert uploader.label == "Carregar o seu questionário"
    assert uploader.proto.max_upload_size_mb == 50
    assert uploader.value is None
    assert any(c.value == "PDF ou DOCX · máximo de 50 MB por ficheiro" for c in a.caption)
    assert any(c.value == EXAMPLE_DESCRIPTION for c in a.caption)
    click_label(a, "Usar questionário de exemplo")
    s = a.session_state["prototype"]
    assert s["stage"] == "structure"
    assert s["instrumento"]["texto_bruto"][0]["texto"] == TEXT
    assert not s["confirmada"] and not s["propostas"]
    assert any(c.value == EXAMPLE_LIMIT for c in a.caption)

    def previous_results():
        # Marcadores exclusivamente técnicos: detetam qualquer resíduo da
        # análise anterior sem executar ou simular uma apreciação metodológica.
        save_representation(s, s["representacao"], True)
        for field in ("verificacoes", "propostas", "decisoes", "acoes"):
            s[field] = {"marcador_teste": {"anterior": True}}
        for field in ("pedidos", "historico"):
            s[field] = [{"marcador_teste": "anterior"}]
        s["documentos"] = [{"tipo": "instruções", "nome": "Documento anterior",
            "texto": "Conteúdo anterior", "confirmado": True,
            "visibilidade": "contexto interno fornecido à aplicação"}]

    def assert_results_cleared():
        assert not s["confirmada"]
        for field in ("verificacoes", "propostas", "decisoes", "acoes", "pedidos", "historico", "documentos"):
            assert not s[field], field

    previous_results()
    old_id = s["id"]
    a.button(key="nav_intake").click().run()
    a.file_uploader[0].set_value(upload)
    a.run()
    # Selecionar o ficheiro só o prepara; a substituição exige o botão do formulário.
    assert s["id"] == old_id
    click_label(a, "Extrair / substituir instrumento")
    assert s["id"] != old_id
    assert s["instrumento"]["nome"] == upload[0]
    assert s["instrumento"]["sha256"] == sha256(content).hexdigest()
    assert s["representacao"] and not s["contexto"]
    assert a.file_uploader[0].value is None
    assert_results_cleared()

    previous_results()
    old_id = s["id"]
    # Um ficheiro ainda no formulário não deve prevalecer sobre o exemplo.
    a.file_uploader[0].set_value(upload)
    next(t for t in a.text_input if t.label == "Designação do questionário").input("Título anterior")
    a.run()
    click_label(a, "Usar questionário de exemplo")
    assert s["id"] != old_id
    assert s["instrumento"]["texto_bruto"][0]["texto"] == TEXT
    assert s["contexto"] and s["stage"] == "structure"
    assert_results_cleared()
    a.button(key="nav_intake").click().run()
    assert not a.exception
    assert a.file_uploader[0].value is None
    assert next(t for t in a.text_input if t.label == "Designação do questionário").value == ""


def test_actual_api_failure_path_does_not_claim_conclusion():
    a = app()
    a.button(key="demo").click().run()
    s = a.session_state["prototype"]
    save_representation(s, s["representacao"], True)
    a.button(key="nav_review").click().run()
    assert not a.exception
    assert any("não executada" in w.value for w in a.warning)
    assert not s["propostas"] and not s["decisoes"]


def test_human_decision_form_requires_explicit_action():
    # Proposta substituta exclusivamente para o teste técnico do formulário.
    from .test_workflow import proposal, fill
    a = app(); a.button(key="demo").click().run()
    s = a.session_state["prototype"]
    save_representation(s, s["representacao"], True)
    fill(s); s["propostas"]["D2.4"] = proposal().model_dump()
    a.button(key="nav_decisions").click().run()
    next(t for t in a.text_input if t.label == "Responsável pela decisão").input("Avaliador fictício do teste técnico")
    next(t for t in a.text_area if t.label == "Fundamentação da decisão humana").input("Validação do percurso de interface com fixture identificada.")
    click_label(a, "Guardar decisão — D2.4")
    assert not s["decisoes"]
    next(c for c in a.checkbox if c.label.startswith("Valido expressamente")).check()
    click_label(a, "Guardar decisão — D2.4")
    assert s["decisoes"]["D2.4"]["expressa"]


def test_request_and_unavailable_response_form():
    from .test_workflow import observation, verification
    from prototipo.models import ReviewBatch
    from prototipo.state import register_batch
    a = app(); a.button(key="demo").click().run()
    s = a.session_state["prototype"]
    save_representation(s, s["representacao"], True)
    obs = observation(informacao_adicional="Qual o destino previsto?", tipo_situacao="esclarecimento necessário", acao="pedido de esclarecimento", estado="A aguardar esclarecimento")
    register_batch(s, ReviewBatch(verificacoes=[verification("D2.4-R01", obs=[obs])]))
    a.button(key="nav_clarifications").click().run()
    next(t for t in a.text_area if t.label == "Resposta ou declaração de indisponibilidade").input("Não disponho desta informação.")
    next(c for c in a.checkbox if c.label == "Não disponho desta informação").check()
    click_label(a, "Guardar resposta / fecho")
    assert s["pedidos"][0]["estado"] == "fechado com limites"
    assert not s["propostas"]
