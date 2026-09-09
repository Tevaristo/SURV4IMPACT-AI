from csv import reader
from io import BytesIO, StringIO
import json
import zipfile

from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from prototipo.models import Element, SELECTION_RULES
from prototipo.norma import ROOT
from prototipo.table_export import (CSV_FILENAME, READING_COLUMNS, SHEET_NAME,
                                    XLSX_FILENAME, export_reading_csv,
                                    export_reading_xlsx, reading_rows)


EXPECTED_COLUMN_GUIDANCE = {
    "Referência": "Identificador único da pergunta, instrução ou bloco. Confirme o código apresentado no questionário original.",
    "Pergunta, instrução ou bloco": "Texto identificado pela aplicação. Confirme se está completo, correto e associado à referência adequada.",
    "Página ou localização no original": "Página, secção, bloco ou outra localização que permita encontrar este elemento no documento original.",
    "Formato de resposta": "Forma prevista para responder: resposta livre, escolha única, escolha múltipla, escala, grelha ou outro formato aplicável.",
    "Opções ou escala de resposta": "Introduza cada opção ou ponto da escala numa linha separada, preservando a ordem e os rótulos do questionário original. Por exemplo: ‘Sim’ na primeira linha e ‘Não’ na segunda.",
    "Subperguntas": "Perguntas ou itens distintos apresentados sob uma pergunta ou enunciado comum. Registe um item por linha.",
    "Linhas da grelha": "Elementos apresentados nas linhas de uma pergunta organizada em grelha ou tabela. Preserve a ordem original.",
    "Colunas da grelha": "Categorias, opções ou pontos da escala apresentados nas colunas de uma grelha ou tabela. Preserve a ordem original.",
    "Instruções aplicáveis": "Orientações aplicáveis à pergunta ou ao bloco, incluindo quem deve responder, como responder e eventuais condições de percurso.",
    "Perguntas a que a instrução se aplica": "Referências das perguntas abrangidas por uma instrução aplicável a várias perguntas. Registe uma referência por linha.",
    "Condições de percurso e destinos": "Condições, filtros ou encaminhamentos e a pergunta ou secção para onde conduzem.",
    "Regra de seleção": "Indicação sobre quantas opções podem ou devem ser selecionadas.",
    "Número mínimo": "Número mínimo de respostas ou opções exigidas, quando estiver explicitamente indicado.",
    "Número máximo": "Número máximo de respostas ou opções permitido, quando estiver explicitamente indicado.",
    "Outra condição de resposta": "Transcrição exata da condição presente no questionário, apenas quando selecionar «Outra condição indicada no questionário».",
}
EXPECTED_HEADER_HELP = {
    "Referência": "Confirme o identificador no questionário original.",
    "Pergunta, instrução ou bloco": "Confirme o texto completo e a referência associada.",
    "Página ou localização no original": "Indique onde encontrar o elemento no original.",
    "Formato de resposta": "Confirme a forma prevista para responder.",
    "Opções ou escala de resposta": "Registe uma opção por linha e preserve a ordem.",
    "Subperguntas": "Registe uma subpergunta por linha.",
    "Linhas da grelha": "Registe uma linha da grelha por linha.",
    "Colunas da grelha": "Registe uma coluna da grelha por linha.",
    "Instruções aplicáveis": "Confirme as orientações aplicáveis ao elemento.",
    "Perguntas a que a instrução se aplica": "Registe uma referência abrangida por linha.",
    "Condições de percurso e destinos": "Confirme as condições e os respetivos destinos.",
    "Regra de seleção": "Escolha a regra explicitamente indicada no original.",
    "Número mínimo": "Introduza um inteiro não negativo, quando aplicável.",
    "Número máximo": "Introduza um inteiro não negativo, quando aplicável.",
    "Outra condição de resposta": "Preencha apenas para outra condição indicada no original.",
}


def export_fixture():
    return [Element(
        id="=1+1",
        texto="Questão sobre satisfação e avaliação\nSegunda linha com ç e ã",
        localizacao="+A1",
        formato="escolha múltipla",
        opcoes=["Ótimo", "Não aplicável"],
        subitens=["-2+3"],
        linhas=["@SOMA(A1:A2)"],
        colunas=["Discordo", "Concordo"],
        instrucoes="Leia com atenção.",
        instrucao_aplica_a=["P1", "P2"],
        condicoes_destinos="Sim → P2\nNão → FIM",
        restricoes="Máximo de duas opções",
    ).model_dump()]


def test_xlsx_is_valid_and_matches_reading_table():
    elements = export_fixture()
    content = export_reading_xlsx(elements)
    assert zipfile.is_zipfile(BytesIO(content))
    workbook = load_workbook(BytesIO(content), data_only=False)
    assert workbook.sheetnames == [SHEET_NAME]
    sheet = workbook[SHEET_NAME]
    headers = [label for _, label in READING_COLUMNS]
    expected = [[row[key] for key, _ in READING_COLUMNS] for row in reading_rows(elements)]
    assert [cell.value for cell in sheet[1]] == headers
    workbook_rows = [["" if cell.value is None else cell.value for cell in row]
                     for row in sheet.iter_rows(min_row=2)]
    assert workbook_rows == expected
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == f"A1:O{len(expected) + 1}"
    assert all(cell.font.bold and cell.fill.fill_type == "solid" for cell in sheet[1])
    assert all(cell.alignment.vertical == "top" and cell.alignment.wrap_text for row in sheet for cell in row)
    assert all(sheet.column_dimensions[column].width > 10 for column in "ABCDEFGHIJKLMNO")
    assert sheet["A2"].value == "=1+1" and sheet["A2"].data_type == "s"
    assert sheet["C2"].value == "+A1" and sheet["C2"].data_type == "s"
    assert "ç e ã" in sheet["B2"].value and "\n" in sheet["B2"].value
    assert "\n" in sheet["E2"].value and sheet["E2"].data_type == "s"
    assert sheet["F2"].value == "-2+3" and sheet["F2"].data_type == "s"
    assert sheet["G2"].value == "@SOMA(A1:A2)" and sheet["G2"].data_type == "s"


def test_csv_remains_functional_and_neutralizes_formulas():
    content = export_reading_csv(export_fixture())
    assert content.startswith(b"\xef\xbb\xbf")
    rows = list(reader(StringIO(content.decode("utf-8-sig")), delimiter=";"))
    assert rows[0] == [label for _, label in READING_COLUMNS]
    assert rows[1][0] == "'=1+1"
    assert rows[1][2] == "'+A1"
    assert rows[1][5] == "'-2+3"
    assert rows[1][6] == "'@SOMA(A1:A2)"
    assert "satisfação" in rows[1][1] and "\n" in rows[1][1]
    assert rows[1][4] == "Ótimo\nNão aplicável"


def test_streamlit_offers_portuguese_xlsx_and_csv_downloads():
    app = AppTest.from_file(str(ROOT / "app_prototipo.py"), default_timeout=30).run()
    app.button(key="demo").click().run()
    downloads = {button.label: button.proto for button in app.download_button}
    assert downloads["Descarregar tabela em Excel"].url.endswith(".xlsx")
    assert downloads["Descarregar tabela em Excel"].type == "primary"
    assert downloads["Descarregar em CSV"].url.endswith(".csv")
    assert downloads["Descarregar em CSV"].type == "secondary"
    other_formats = next(panel for panel in app.expander if panel.label == "Outros formatos")
    assert not other_formats.proto.expanded
    assert [button.label for button in other_formats.download_button] == ["Descarregar em CSV"]
    columns = json.loads(app.dataframe[0].proto.columns)
    assert [columns[key]["label"] for key, _ in READING_COLUMNS] == [label for _, label in READING_COLUMNS]
    assert columns["regra_selecao"]["required"] is True
    assert columns["regra_selecao"]["type_config"]["options"] == SELECTION_RULES
    for key in ("numero_minimo", "numero_maximo"):
        assert columns[key]["type_config"]["type"] == "number"
        assert columns[key]["type_config"]["min_value"] == 0
        assert columns[key]["type_config"]["step"] == 1


def test_reading_table_has_general_and_per_column_guidance():
    app = AppTest.from_file(str(ROOT / "app_prototipo.py"), default_timeout=30).run()
    app.button(key="demo").click().run()
    general = ("Compare cada linha com o questionário original. Corrija os elementos que não tenham sido identificados "
               "corretamente e assinale como não confirmada a informação que não consiga verificar.")
    assert any(node.value == general for node in app.markdown)
    panel = next(item for item in app.expander if item.label == "Como rever esta tabela")
    assert not panel.proto.expanded
    guidance = "\n".join(node.value for node in panel.markdown)
    assert list(EXPECTED_COLUMN_GUIDANCE) == [label for _, label in READING_COLUMNS]
    for label, instruction in EXPECTED_COLUMN_GUIDANCE.items():
        assert f"**{label}:** {instruction}" in guidance
    columns = json.loads(app.dataframe[0].proto.columns)
    for key, label in READING_COLUMNS:
        assert columns[key]["label"] == label
        assert columns[key]["help"] == EXPECTED_HEADER_HELP[label]
        assert len(columns[key]["help"]) < len(EXPECTED_COLUMN_GUIDANCE[label])


def test_multiline_option_detail_is_collapsed_editable_and_preserves_line_breaks():
    app = AppTest.from_file(str(ROOT / "app_prototipo.py"), default_timeout=30).run()
    app.button(key="demo").click().run()
    panel = next(item for item in app.expander if item.label == "Ver ou editar opções e escalas por pergunta")
    assert not panel.proto.expanded
    assert panel.selectbox[0].value == "P2a"
    assert panel.text_area[0].value == "1\n2\n3\n4\n5"
    panel.text_area[0].input("1\n2\n3")
    panel.button[0].click().run()
    element = next(value for value in app.session_state["prototype"]["representacao"] if value["id"] == "P2a")
    assert element["opcoes"] == ["1", "2", "3"]
    assert not app.session_state["prototype"]["confirmada"]


def test_demo_table_displays_lists_with_preserved_line_breaks():
    app = AppTest.from_file(str(ROOT / "app_prototipo.py"), default_timeout=30).run()
    app.button(key="demo").click().run()
    table = app.dataframe[0].value.set_index("id")
    assert table.loc["P2a", "opcoes"] == "1\n2\n3\n4\n5"
    assert table.loc["P6.2", "opcoes"] == "Sim\nNão"
    assert table.loc["P6.4", "opcoes"].splitlines()[1] == "Sim, mas com alterações: menor dimensão financeira; redução no valor de investimento em %"
    assert len(table.loc["P7", "opcoes"].splitlines()) == 4
    assert table.loc["P7", "subitens"] == ""
