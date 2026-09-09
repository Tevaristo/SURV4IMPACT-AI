import json
import sys
from pathlib import Path

from docx import Document


def clean(text: str) -> str:
    return " ".join(text.replace("\xa0", " ").split())


def extract(path: Path) -> dict:
    document = Document(path)
    paragraphs = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = clean(paragraph.text)
        if text:
            paragraphs.append(
                {
                    "index": index,
                    "style": paragraph.style.name if paragraph.style else "",
                    "text": text,
                }
            )

    tables = []
    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            rows.append([clean(cell.text) for cell in row.cells])
        tables.append({"index": table_index, "rows": rows})

    return {
        "source": str(path),
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "paragraphs": paragraphs,
        "tables": tables,
    }


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: extract_docx_structure.py INPUT.docx OUTPUT.json")
    source = Path(sys.argv[1])
    destination = Path(sys.argv[2])
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(
        json.dumps(extract(source), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
